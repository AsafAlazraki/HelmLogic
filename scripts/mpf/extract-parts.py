#!/usr/bin/env python3
"""MPF Phase-2 PARTS extraction (STRICTLY READ-ONLY on the xlsx sources).

Extracts five datasets to tasks/mpf-audit/extracted/:
  dealer-fit.json         Parts Module.xlsx  / 'Dealer Fit Module'  (1,795 DFOs)
  parts-maintenance.json  Parts Module.xlsx  / 'Parts Maintenance'  (3,667 keyed rows, ~135 dup codes)
  parts-inventory.json    Parts Module.xlsx  / 'Parts Data Drop'    (26,347 rows, key Franchise+Part)
  rigging-kits.json       Rigging Module.xlsx/ 'Rigging Kits'       (961 PNs; NLA / #N/A quarantined)
  suppliers.json          Supplier Module.xlsx / 'Sheet1'           (1,606 suppliers)

Per the phase-1 analysis (tasks/mpf-audit/analysis/*.md):
- bound-scan: iterate real rows, never trust max_row; skip section pseudo-header rows
- keys trimmed (DFO ' 9HI_HBP 171_PD' anomaly); numeric part numbers cast to str
- Excel error artifacts -> null / quarantine
- GST: only Parts Data Drop 'Retail+ GST' is inc-GST in this wave -> ex-GST derived,
  normalization recorded per row. All other master-sheet prices are ex GST.
- D10: masters only — the 15 vendor price-list sheets are NOT extracted this wave.
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _common import (SOURCE_DIR, EXCEL_ERRORS, append_audit, clean_number, clean_text,
                     ex_gst, is_excel_error, money, slug, write_dataset)

import openpyxl

PARTS_XLSX = os.path.join(SOURCE_DIR, "Parts Module.xlsx")
RIGGING_XLSX = os.path.join(SOURCE_DIR, "Rigging Module.xlsx")
SUPPLIER_XLSX = os.path.join(SOURCE_DIR, "Supplier Module.xlsx")


# ------------------------------------------------------------------ Dealer Fit Module

def extract_dealer_fit(wb):
    """Header row 11, data from row 12. Rollup cols 2-18, 30 accessory slots of 9
    cols starting col 19 (slot 1 carries Image Link at col 27), Long Description col 290."""
    ws = wb["Dealer Fit Module"]
    rows, sections = [], []
    seen_keys = {}
    section = None
    n_section_rows = 0
    for i, r in enumerate(ws.iter_rows(min_row=12, values_only=True), 12):
        desc = clean_text(r[2]) if len(r) > 2 else None
        raw_key = r[3] if len(r) > 3 else None
        key = clean_text(raw_key)
        if key is None:
            if desc is not None:
                n_section_rows += 1
                section = desc
                sections.append({"name": desc, "sourceRow": i})
            continue
        # trimmed natural key (leading/trailing whitespace anomaly §6.3)
        key = str(raw_key).strip()
        components = []
        for k in range(30):
            base = 19 + 9 * k
            if base >= len(r):
                break
            c_name = clean_text(r[base])
            c_code = clean_text(r[base + 1]) if base + 1 < len(r) else None
            if c_name is None and c_code is None:
                continue
            comp = {
                "slot": k + 1,
                "name": c_name,
                "code": c_code,
                "ctd": money(r[base + 2]) if base + 2 < len(r) else None,
                "sell": money(r[base + 3]) if base + 3 < len(r) else None,
                "labour": money(r[base + 4]) if base + 4 < len(r) else None,
                "labHrs": clean_number(r[base + 5], 3) if base + 5 < len(r) else None,
                "sundry": money(r[base + 6]) if base + 6 < len(r) else None,
                "sublet": money(r[base + 7]) if base + 7 < len(r) else None,
            }
            components.append(comp)
        row = {
            "key": key,
            "keyRaw": str(raw_key),
            "keyTrimmed": key != str(raw_key),
            "docId": slug(key),
            "name": desc,
            "section": section,
            "ctd": money(r[4]),
            "adjCtd": money(r[6]),
            "totPartsCtd": money(r[7]),
            "partsSell": money(r[8]),
            "totalLabHrs": clean_number(r[9], 3),
            "labourCtd": money(r[10]),
            "labourRet": money(r[11]),
            "sundry": money(r[12]),
            "sublet": money(r[13]),
            "actCtd": money(r[14]),
            "actSell": money(r[17]),
            "rebate": money(r[18]),
            "imageLink": clean_text(r[27]) if len(r) > 27 else None,
            "longDescription": clean_text(r[290]) if len(r) > 290 else None,
            "components": components,
            "gstBasis": "exGst",  # analysis §1: master-sheet sells/costs are ex GST
            "sourceRow": i,
        }
        seen_keys.setdefault(key, []).append(i)
        rows.append(row)

    dups = {k: v for k, v in seen_keys.items() if len(v) > 1}
    meta = {
        "source": "tasks/mpf-source/Parts Module.xlsx#'Dealer Fit Module'",
        "target": "organisations/{orgId}/dealerFitSelections (doc id = slug of trimmed DFO key)",
        "sectionRows": n_section_rows,
        "distinctKeys": len(seen_keys),
        "duplicateKeys": dups,
        "trimmedKeys": [r["key"] for r in rows if r["keyTrimmed"]],
        "gstNote": "All rollup prices ex GST (Act CTD / Act Sell convention). No conversion applied.",
    }
    write_dataset("dealer-fit.json", meta, rows, sections=sections)
    return {"rows": len(rows), "sections": n_section_rows, "distinctKeys": len(seen_keys),
            "duplicateKeys": len(dups), "trimmedKeys": len(meta["trimmedKeys"])}


# ------------------------------------------------------------------ Parts Maintenance

def extract_parts_maintenance(wb):
    """Header row 1, row 2 = column numbers, data from row 3. ~135 duplicate codes
    -> composite key (code / code~2 / code~3 ...) + provenance per occurrence."""
    ws = wb["Parts Maintenance"]
    rows = []
    seen = {}
    section = None
    n_section_rows = 0
    n_no_key_with_prices = 0
    for i, r in enumerate(ws.iter_rows(min_row=3, values_only=True), 3):
        display = clean_text(r[2]) if len(r) > 2 else None
        code = clean_text(r[4]) if len(r) > 4 else None
        if code is None:
            if display is not None:
                n_section_rows += 1
                section = display
                if money(r[11]) is not None or money(r[8]) is not None:
                    n_no_key_with_prices += 1
            continue
        code = str(r[4]).strip()
        occ = seen.get(code, 0) + 1
        seen[code] = occ
        composite = code if occ == 1 else f"{code}~{occ}"
        rows.append({
            "code": code,
            "compositeKey": composite,
            "docId": slug(composite),
            "duplicateIndex": occ,
            "desc": display,
            "supplierDesc": clean_text(r[5]),
            "supplier": clean_text(r[3]),
            "section": section,
            "paCost": money(r[6]),
            "ctd": money(r[8]),
            "sell": money(r[11]),
            "baseList": money(r[12]),
            "installType": clean_text(r[13]),
            "installHrs": clean_number(r[14], 3),
            "labourDollars": money(r[15]),
            "partsCtd": money(r[16]),
            "sundryCtd": money(r[17]),
            "totalCtd": money(r[18]),
            "sellIncInstall": money(r[24]) if len(r) > 24 else None,
            "opCode": clean_text(r[26]) if len(r) > 26 else None,
            "opDesc": clean_text(r[27]) if len(r) > 27 else None,
            "gstBasis": "exGst",
            "sourceRow": i,
        })
    dups = {k: v for k, v in seen.items() if v > 1}
    meta = {
        "source": "tasks/mpf-source/Parts Module.xlsx#'Parts Maintenance'",
        "target": "organisations/{orgId}/fitUpItems (doc id = slug of composite key)",
        "sectionRows": n_section_rows,
        "distinctCodes": len(seen),
        "duplicateCodes": dups,
        "duplicateRowCount": sum(v - 1 for v in dups.values()),
        "noKeyRowsWithPrices": n_no_key_with_prices,
        "gstNote": "CTD / Sell / Sell inc Install are ex GST per analysis. No conversion applied.",
    }
    write_dataset("parts-maintenance.json", meta, rows)
    return {"rows": len(rows), "sections": n_section_rows, "distinctCodes": len(seen),
            "duplicateCodes": len(dups), "duplicateRows": meta["duplicateRowCount"]}


# ------------------------------------------------------------------ Parts Data Drop

def extract_parts_inventory(wb):
    """Header row 1, data from row 2 (first data at row 3). Key = Franchise + Part.
    'Retail+ GST' is inc GST -> ex-GST derived and recorded per row."""
    ws = wb["Parts Data Drop"]
    rows = []
    seen = {}
    for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        franchise = clean_text(r[2]) if len(r) > 2 else None
        part_raw = r[3] if len(r) > 3 else None
        part = clean_text(part_raw)
        if franchise is None or part is None:
            continue
        # numbers-as-part-numbers: cast without float artifacts (1000.0 -> '1000')
        if isinstance(part_raw, float) and part_raw.is_integer():
            part = str(int(part_raw))
        retail_inc = money(r[11])
        retail_ex = ex_gst(retail_inc)
        key = f"{franchise}::{part}"
        occ = seen.get(key, 0) + 1
        seen[key] = occ
        rows.append({
            "franchise": franchise,
            "part": part,
            "key": key,
            "docId": slug(f"{franchise}-{part}") if occ == 1 else slug(f"{franchise}-{part}") + f"--{occ}",
            "desc": clean_text(r[4]),
            "stockOH": clean_number(r[5], 3),
            "bin": clean_text(r[6]),
            "cost": money(r[7]),          # 'Daily' = daily-average cost
            "list": money(r[10]),         # ex GST
            "retailIncGst": retail_inc,   # source value (inc GST)
            "retailExGst": retail_ex,     # derived
            "gstNormalized": ({"retailExGst": "Retail+ GST / 1.1"} if retail_inc is not None else None),
            "sourceRow": i,
        })
    dups = {k: v for k, v in seen.items() if v > 1}
    meta = {
        "source": "tasks/mpf-source/Parts Module.xlsx#'Parts Data Drop' (DMS extract as at 27.02.2026)",
        "target": "organisations/{orgId}/serviceParts (doc id = slug of franchise+part)",
        "distinctKeys": len(seen),
        "duplicateKeys": dups,
        "gstNote": "'Retail+ GST' inc GST -> retailExGst derived (/1.1), recorded per row. List / Daily are ex GST.",
    }
    write_dataset("parts-inventory.json", meta, rows)
    return {"rows": len(rows), "distinctKeys": len(seen), "duplicateKeys": len(dups),
            "gstNormalizedRows": sum(1 for x in rows if x["gstNormalized"])}


# ------------------------------------------------------------------ Rigging Kits

RK_KNOWN_BUILDS = {"Service", "Factory", "Base List"}


def extract_rigging_kits():
    """Headers row 1, constants row 2, column numbers row 3, data from row 4.
    26 section pseudo-headers inside the data range. NLA-in-key + #N/A price rows quarantined."""
    wb = openpyxl.load_workbook(RIGGING_XLSX, read_only=True, data_only=True)
    ws = wb["Rigging Kits"]
    rows, quarantined, sections = [], [], []
    seen = {}
    section = None
    for i, r in enumerate(ws.iter_rows(min_row=4, values_only=True), 4):
        desc = clean_text(r[2]) if len(r) > 2 else None
        pn_raw = r[3] if len(r) > 3 else None
        pn = clean_text(pn_raw)
        if pn is None:
            if desc is not None:
                section = desc
                sections.append({"name": desc, "sourceRow": i})
            continue
        if isinstance(pn_raw, (int, float)):  # anomaly §6.4 numeric PN (7851816805)
            pn = str(int(pn_raw)) if float(pn_raw).is_integer() else str(pn_raw)
        pn = pn.strip()

        # NLA embedded in the key (anomaly §6.1)
        status = "active"
        clean_pn = pn
        if "NLA" in pn.upper():
            status = "nla"
            clean_pn = pn.upper().replace("###", " ").replace("NLA", " ")
            clean_pn = " ".join(clean_pn.split()).strip()
            # keep original casing for the non-NLA part
            import re as _re
            clean_pn = _re.sub(r"\s*#*\s*NLA\s*#*\s*", " ", pn, flags=_re.I).strip()
            clean_pn = " ".join(clean_pn.replace("#", " ").split())

        price_cells = [r[5], r[6], r[7], r[10], r[11], r[12]]
        has_error_prices = any(is_excel_error(c) for c in price_cells)

        build_raw = clean_text(r[4])
        build = build_raw if build_raw in RK_KNOWN_BUILDS else None

        components = []
        for base in (32, 34, 36, 38):
            c_name = clean_text(r[base]) if base < len(r) else None
            if c_name is None:
                continue
            components.append({"name": c_name, "ctd": money(r[base + 1]) if base + 1 < len(r) else None})

        control_cables = []
        for c in range(45, 55):
            v = clean_text(r[c]) if c < len(r) else None
            if v is not None:
                control_cables.append(v)

        row = {
            "partNo": clean_pn,
            "partNoRaw": pn,
            "docId": slug(clean_pn),
            "status": status,
            "description": desc,
            "section": section,
            "build": build,
            "buildRaw": (build_raw if build_raw != build else None),  # anomaly §6.6 pasted prices
            "dealerCost": money(r[5]),
            "freight": money(r[6]),          # header says 'Factory'; content is freight (anomaly §6.5)
            "kitCtd": money(r[7]),
            "retailExGst": money(r[10]),
            "tradeExGst": money(r[11]),
            "subDealerExGst": money(r[12]),
            "installHrs": clean_number(r[14], 3),
            "installLabour": money(r[15]),
            "installAdditionalParts": money(r[16]),
            "installSundry": money(r[17]),
            "totalCtd": money(r[25]),
            "totalSellExGst": money(r[28]),
            "totalTradeExGst": money(r[29]),
            "totalSubDealerExGst": money(r[30]),
            "components": components,
            "inclusionFlags": {
                "fuelFilter": clean_text(r[41]),
                "controlCables": control_cables,
            },
            "dealerCost2022": money(r[56]) if len(r) > 56 else None,
            "gstBasis": "exGst",
            "sourceRow": i,
        }
        if status == "nla" or has_error_prices:
            row["quarantineReason"] = "; ".join(
                x for x in [
                    "NLA marker embedded in part number" if status == "nla" else None,
                    "#N/A / error artifact in price columns" if has_error_prices else None,
                ] if x)
            quarantined.append(row)
            continue
        seen.setdefault(clean_pn, []).append(i)
        rows.append(row)
    wb.close()

    dups = {k: v for k, v in seen.items() if len(v) > 1}
    meta = {
        "source": "tasks/mpf-source/Rigging Module.xlsx#'Rigging Kits'",
        "target": "organisations/{orgId}/riggingKits (NEW collection per D8; doc id = slug of partNo)",
        "sections": len(sections),
        "distinctPartNos": len(seen),
        "duplicatePartNos": dups,
        "quarantinedCount": len(quarantined),
        "rateConstants": {"retailMu": 0.25, "tradeMu": 0.05, "subDealerMu": 0.05, "labourRateExGst": 130.09},
        "gstNote": "All three tiers + install pricing ex GST per analysis. No conversion applied.",
    }
    write_dataset("rigging-kits.json", meta, rows, quarantined=quarantined, sections=sections)
    return {"rows": len(rows), "quarantined": len(quarantined), "sections": len(sections),
            "distinctPartNos": len(seen), "duplicatePartNos": len(dups)}


# ------------------------------------------------------------------ Suppliers

def extract_suppliers():
    wb = openpyxl.load_workbook(SUPPLIER_XLSX, read_only=True, data_only=True)
    ws = wb["Sheet1"]
    rows = []
    seen = {}
    for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        sid = clean_text(r[0])
        if sid is None:
            continue
        sid = str(r[0]).strip()
        seen.setdefault(sid, []).append(i)
        terms = clean_text(r[7])
        if terms == "<None>":
            terms = None
        rows.append({
            "supplierId": sid,
            "docId": slug(sid),
            "name": clean_text(r[1]),
            "address": {
                "address1": clean_text(r[2]),
                "suburb": clean_text(r[3]),
                "postCode": (str(r[4]).strip() if clean_text(r[4]) is not None else None),
                "state": clean_text(r[5]),
            },
            "phone": clean_text(r[6]),
            "paymentTerms": terms,
            "email": clean_text(r[8]),
            "contact": clean_text(r[9]),
            "abn": clean_text(r[10]),
            "acn": clean_text(r[11]),
            "mobile": clean_text(r[12]),
            "creditLimit": money(r[13]),
            "paymentType": clean_text(r[14]),
            "dmsConfig": {
                "genReceiptInvoiceFormat": clean_text(r[15]),
                "genReceiptInvoiceTaxType": clean_text(r[16]),
                "partsReceiptInvoiceFormat": clean_text(r[17]),
                "partsReceiptGstCalc": clean_text(r[18]),
            },
            "sourceRow": i,
        })
    wb.close()
    dups = {k: v for k, v in seen.items() if len(v) > 1}
    meta = {
        "source": "tasks/mpf-source/Supplier Module.xlsx#'Sheet1'",
        "target": "organisations/{orgId}/suppliers (NEW collection; doc id = slug of Supplier Id)",
        "distinctIds": len(seen),
        "duplicateIds": dups,
        "gstNote": "No price columns; no conversion.",
    }
    write_dataset("suppliers.json", meta, rows)
    return {"rows": len(rows), "distinctIds": len(seen), "duplicateIds": len(dups)}


def main():
    print("=== MPF Phase-2 PARTS extraction (sources read-only) ===")
    summary = {}
    print("\n[1/5+2/5+3/5] Parts Module.xlsx (40.5 MB, streaming)...")
    wb = openpyxl.load_workbook(PARTS_XLSX, read_only=True, data_only=True)
    summary["dealerFit"] = extract_dealer_fit(wb)
    summary["partsMaintenance"] = extract_parts_maintenance(wb)
    summary["partsInventory"] = extract_parts_inventory(wb)
    wb.close()
    print("\n[4/5] Rigging Module.xlsx...")
    summary["riggingKits"] = extract_rigging_kits()
    print("\n[5/5] Supplier Module.xlsx...")
    summary["suppliers"] = extract_suppliers()
    print("\n=== summary ===")
    for k, v in summary.items():
        print(f"  {k}: {v}")
    return summary


if __name__ == "__main__":
    main()
