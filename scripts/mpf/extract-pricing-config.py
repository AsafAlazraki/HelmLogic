#!/usr/bin/env python3
"""MPF Phase 2 — Pricing / Rego / Freight config extraction (STRICTLY READ-ONLY on sources).

Sources:
  tasks/mpf-source/Price Matrix.xlsx        -> pricing-matrix.json + exchange-rates.json
  tasks/mpf-source/Registration Module.xlsx -> rego-catalog.json (state: QLD explicit)
  tasks/mpf-source/Freight Module.xlsx      -> freight-config.json (per-vendor $/lm + buffer)

Outputs to tasks/mpf-audit/extracted/.
"""
import json, os, datetime
from openpyxl import load_workbook

OUT = "tasks/mpf-audit/extracted"
os.makedirs(OUT, exist_ok=True)
NBSP = "\xa0"
NOW = datetime.datetime.now(datetime.timezone.utc).isoformat()

CURRENCY_MAP = {"$A": "AUD", "$NZ": "NZD", "Euro": "EUR", "EURO": "EUR", "$US": "USD"}


def clean(v):
    if v is None:
        return None
    if isinstance(v, str):
        s = v.replace(NBSP, " ").strip()
        return s or None
    return v


def cell(v):
    """margin cell -> float fraction | 'RRP' | None"""
    v = clean(v)
    if v is None:
        return None
    if isinstance(v, str):
        return "RRP" if v.upper() == "RRP" else v
    return round(float(v), 4)


def dt(v):
    v = clean(v)
    if v is None:
        return None
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    return str(v)


# ───────────────────────── Price Matrix ─────────────────────────
wb = load_workbook("tasks/mpf-source/Price Matrix.xlsx", read_only=True, data_only=True)
ws = wb["Price Matrix"]
rows = list(ws.iter_rows(min_row=1, max_row=70, max_col=21, values_only=True))

REGIONS = {  # sheet row -> region label
    **{r: "boat-brands" for r in range(8, 15)},
    **{r: "trailers" for r in range(16, 22)},
    **{r: "motors" for r in range(23, 26)},
    **{r: "accessories-electronics" for r in range(27, 56)},
    **{r: "sublets-admin" for r in range(58, 60)},
}
franchises = []
for rnum in sorted(REGIONS):
    row = rows[rnum - 1]
    brand, code = clean(row[2]), clean(row[3])
    if brand is None:
        continue
    fx = clean(row[5])
    sell = cell(row[13])
    franchises.append({
        "brand": brand,
        "franchiseCode": code,
        "region": REGIONS[rnum],
        "buyCurrency": CURRENCY_MAP.get(fx, fx),
        "buyCurrencyRaw": fx,
        "bmtLabourRate": clean(row[6]),      # 'Internal'
        "retailLabourRate": clean(row[7]),   # 'Retail'
        "reviewed": dt(row[8]),
        "ctdLoad": cell(row[9]),
        "other": cell(row[10]),
        "subDealerDiscount": cell(row[11]),   # discount OFF retail (0 = none; negative = charged above base)
        "tradeDiscount": cell(row[12]),
        "sellMarkup": sell,                   # margin applied on cost to build retail; 'RRP' = sell at supplier RRP
        "factoryOptionsMarkup": cell(row[14]),
        "dealerFitMarkup": cell(row[15]),
        "warrantyAllowance": cell(row[17]),
        "adminLoad": cell(row[18]),
        "isRrp": sell == "RRP",
        "notes": clean(row[20]),
        "sourceRow": rnum,
    })

# Retail Sliding Scale rows 62-69 — band bounds are fixed knowledge from the labels
SCALE_BOUNDS = [(0, 10.0), (10.01, 25.0), (25.01, 35.0), (35.01, 50.0),
                (50.01, 100.0), (100.01, 500.0), (500.01, 1000.0), (1000.01, 2500.0)]
sliding_scale = []
for k, rnum in enumerate(range(62, 70)):
    row = rows[rnum - 1]
    label = clean(row[2])
    if label is None:
        continue
    lo, hi = SCALE_BOUNDS[k]
    sliding_scale.append({
        "label": label, "minPrice": lo, "maxPrice": hi,
        "margin": cell(row[10]),  # mirrored in Other + Sell cols
        "sourceRow": rnum,
    })

pm_out = {
    "meta": {
        "source": "tasks/mpf-source/Price Matrix.xlsx", "sheet": "Price Matrix",
        "extractedUtc": NOW, "franchiseCount": len(franchises),
        "bannerNotes": [
            "NB: Trade/Sub Dealer Pricing is Discount off Retail",
            "NB: All Admin Loads removed 18.08.2025 as per MM / JF Request",
        ],
        "semantics": {
            "valuesAreFractions": "0.29 = 29%",
            "sellMarkup": "margin on cost to build retail; RRP = sell at supplier RRP (no computed margin)",
            "subDealerDiscount/tradeDiscount": "discount OFF retail per banner; NEGATIVE = charged above base (Rigging Kits -0.05)",
            "otherColumn": "duplicates Sell on nearly all rows — legacy, semantics undocumented",
        },
        "knownCollisions": [
            "Stacer boats and Stacer Trailers share franchise code 9ST with different margin rows",
            "Jeanneau and Merry Fisher share 9JE",
            "Yamaha motors and Rigging Kits share 9YA",
            "Price Matrix 9SC (Stabicraft) vs Boat CSV PD-code prefix 9SB",
        ],
    },
    "franchises": franchises,
    "retailSlidingScale": sliding_scale,
}
with open(f"{OUT}/pricing-matrix.json", "w") as f:
    json.dump(pm_out, f, indent=1)
print(f"pricing-matrix.json: {len(franchises)} franchise rows + {len(sliding_scale)} sliding-scale bands")

# ───────────────────────── Exchange Rates ─────────────────────────
ws = wb["Exchange Rates"]
fx_rows = []
for row in ws.iter_rows(min_row=10, max_row=15, max_col=8, values_only=True):
    cur = clean(row[2])
    if cur is None:
        continue
    code = {"NZ": "NZD", "EURO": "EUR"}.get(cur, cur)
    fx_rows.append({
        "currency": cur, "code": code,
        "reviewDate": dt(row[3]),
        "rate": float(clean(row[5])),
        "notes": clean(row[7]),
    })
fx_out = {
    "meta": {
        "source": "tasks/mpf-source/Price Matrix.xlsx", "sheet": "Exchange Rates",
        "extractedUtc": NOW,
        "direction": "divisor-style buy rates: AUD = foreign / rate (matches HelmLogic convention — highfield-pricing-workspace computes baseAud = totalUsd / exchangeRate)",
    },
    "rates": fx_rows,
}
with open(f"{OUT}/exchange-rates.json", "w") as f:
    json.dump(fx_out, f, indent=1)
print(f"exchange-rates.json: {len(fx_rows)} rates -> " +
      ", ".join(f"{r['code']}={r['rate']}" for r in fx_rows))
wb.close()

# ───────────────────────── Registration (QLD) ─────────────────────────
wb = load_workbook("tasks/mpf-source/Registration Module.xlsx", read_only=True, data_only=True)
ws = wb["Registration Costs"]
reg_rows = {i: r for i, r in enumerate(
    ws.iter_rows(min_row=1, max_row=34, max_col=12, values_only=True), 1)}


def reg(i, group, applies_to, concession=False, min_len=None, max_len=None,
        min_wt=None, max_wt=None):
    row = reg_rows[i]
    item = {
        "name": clean(row[2]), "group": group, "appliesTo": applies_to,
        "revCode": clean(row[6]),
        "ctd": round(float(clean(row[9])), 2), "sell": round(float(clean(row[10])), 2),
        "concession": concession, "sourceRow": i,
    }
    if min_len is not None:
        item["minLengthM"], item["maxLengthM"] = min_len, max_len
    if min_wt is not None:
        item["minWeightT"], item["maxWeightT"] = min_wt, max_wt
    return item


items = [
    reg(9, "Boat Registration", "boat", min_len=0, max_len=4.5),
    reg(10, "Boat Registration", "boat", min_len=4.51, max_len=6.0),
    reg(11, "Boat Registration", "boat", min_len=6.01, max_len=10.0),
    reg(12, "Boat Registration", "boat", min_len=10.01, max_len=15.0),
    reg(13, "Boat Registration", "boat"),                    # Not Required, $0
    reg(16, "Trailer Registration", "trailer", min_wt=0, max_wt=1.02),
    reg(17, "Trailer Registration", "trailer", min_wt=1.021, max_wt=4.55),
    reg(18, "Trailer Registration", "trailer", min_wt=4.551, max_wt=99),  # Heavy, no REV code
    reg(19, "Trailer Registration", "trailer"),              # Not Required, $0
    reg(22, "Other Fees & Charges", "fee"),                  # Boat Transfer  REGO 5
    reg(23, "Other Fees & Charges", "fee"),                  # Trailer Transfer REGO 8
    reg(24, "Other Fees & Charges", "fee"),                  # Replacement Plate
    reg(25, "Other Fees & Charges", "fee"),                  # Unregistered Vehicle Permit
    reg(26, "Other Fees & Charges", "fee"),                  # VIN Plate
    reg(27, "Other Fees & Charges", "fee"),                  # PPSR Fee
    reg(30, "Boat Registration - Pensioner / Concession", "boat", True, 0, 4.5),
    reg(31, "Boat Registration - Pensioner / Concession", "boat", True, 4.51, 6.0),
    reg(32, "Boat Registration - Pensioner / Concession", "boat", True, 6.01, 10.0),
    reg(33, "Boat Registration - Pensioner / Concession", "boat", True, 10.01, 15.0),
]
rego_out = {
    "meta": {
        "source": "tasks/mpf-source/Registration Module.xlsx", "sheet": "Registration Costs",
        "extractedUtc": NOW,
        "state": "QLD",  # explicit — sheet is single-state with no state column
        "asAt": "2025-07-01",
        "count": len(items),
        "notes": [
            "SELL = CTD rounded UP to whole dollars (convention, not GST — rego fees are GST-free)",
            "Boat rego banded by hull length; trailer rego banded by weight class (tonnes)",
            "Heavy Trailers band and concession bands have no REV code in source",
            "Rego prices change every July — effective-date handling is an open decision",
        ],
    },
    "state": "QLD",
    "items": items,
}
with open(f"{OUT}/rego-catalog.json", "w") as f:
    json.dump(rego_out, f, indent=1)
print(f"rego-catalog.json: {len(items)} QLD items "
      f"({sum(1 for x in items if x['revCode'])} with REV codes, "
      f"{sum(1 for x in items if x['concession'])} concession)")
wb.close()

# ───────────────────────── Freight ─────────────────────────
wb = load_workbook("tasks/mpf-source/Freight Module.xlsx", read_only=True, data_only=True)


def freight_sheet(sheet, vendor_key):
    ws = wb[sheet]
    g = {i: r for i, r in enumerate(
        ws.iter_rows(min_row=1, max_row=40, max_col=9, values_only=True), 1)}
    local_charges = []
    for i in range(15, 28):
        name = clean(g[i][2])
        if name is None or "Sub Total" in str(name):
            continue
        local_charges.append({
            "name": name, "currency": CURRENCY_MAP.get(clean(g[i][4]), clean(g[i][4])),
            "amount": float(clean(g[i][5])), "exRate": float(clean(g[i][7])),
            "aud": round(float(clean(g[i][8])), 2),
        })
    sea = g[10]
    per_lm_row = g[35]
    return {
        "vendorKey": vendor_key,
        "supplier": clean(g[2][2]),
        "contact": clean(g[5][3]),
        "quoteDate": dt(g[6][3]),
        "shipmentRef": clean(g[7][3]),
        "originNote": clean(g[9][2]),
        "seafreight": {
            "name": clean(sea[2]), "currency": CURRENCY_MAP.get(clean(sea[4]), clean(sea[4])),
            "amount": float(clean(sea[5])), "exRate": float(clean(sea[7])),
            "aud": round(float(clean(sea[8])), 2),
        },
        "localCharges": local_charges,
        "localChargesSubtotalAud": round(sum(c["aud"] for c in local_charges), 2),
        "estimateTotalAud": round(float(clean(g[31][8])), 2),
        "bufferPct": float(clean(g[32][3])),
        "seafreightCtdAud": round(float(clean(g[34][8])), 2),
        "sampleContainerLinearMetres": float(clean(per_lm_row[7])),
        "perLinearMetreAud": round(float(clean(per_lm_row[8])), 2),
    }


vendors = [
    freight_sheet("FCL Import - Highfield", "highfield"),
    freight_sheet("Quadrant Pacific - Surtees", "surtees"),
]
freight_out = {
    "meta": {
        "source": "tasks/mpf-source/Freight Module.xlsx",
        "sheets": ["FCL Import - Highfield", "Quadrant Pacific - Surtees",
                   "Freight Distribution Calculator (operational tooling — not extracted as config)"],
        "extractedUtc": NOW,
        "notes": [
            "per-lm rates differ 3.4x between suppliers (Highfield $128.47/lm vs Surtees $441.21/lm); buffers differ (10% vs 5%) — config MUST be per-vendor",
            "Distribution Calculator allocates a container's landed freight across boats by linear metre with small-boat/loaded-boat split — per-shipment ops, not catalog",
        ],
    },
    "vendors": vendors,
}
with open(f"{OUT}/freight-config.json", "w") as f:
    json.dump(freight_out, f, indent=1)
print("freight-config.json: " + ", ".join(
    f"{v['vendorKey']} ${v['perLinearMetreAud']}/lm buffer {v['bufferPct']:.0%}" for v in vendors))
wb.close()
