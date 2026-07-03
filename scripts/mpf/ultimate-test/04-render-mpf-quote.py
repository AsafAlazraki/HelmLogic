#!/usr/bin/env python3
"""Phase 6 Ultimate Test — step 4: render the MPF quote row as a visual.

Produces the "their system" exhibit: the ACTUAL Boat Module sheet, reduced
to the quote-relevant columns of the SP560 (PVC) W-W-WB row (829), printed
to PDF via LibreOffice, then rasterised to PNG via pypdfium2.

Method (presentation only — values untouched):
  1. Load the /tmp/ultimate copy with data_only=True (cached values).
  2. HIDE every row except headers (1-3), the Highfield brand divider (278)
     and the quote row (829); hide every column not in the exhibit set.
  3. Set print area + landscape + fit-to-one-page, save a NEW render copy.
  4. soffice --convert-to pdf  ->  pdftoppm-equivalent via pypdfium2.

Outputs: /tmp/ultimate/render/mpf-quote.pdf + mpf-quote.png
"""
import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
import os, subprocess, sys

UD = "/tmp/ultimate"
SRC = f"{UD}/Boat Module.xlsx"
RENDER_DIR = f"{UD}/render"
RENDER_XLSX = f"{RENDER_DIR}/mpf-quote.xlsx"
ROW = 829
HIGHFIELD_DIVIDER_ROW = 278  # brand divider carrying Highfield-specific header labels

# Exhibit columns (1-based) — identity, rego, motor slot 1, trailer,
# dealer-fit lines 1-2, landed cost, hull Cash price (the quote figures).
KEEP_COLS = [3, 4, 259, 299, 308, 309, 310, 311, 312, 313, 314, 315,
             390, 402, 403, 460, 461]
LAST_COL = 480  # print range upper bound (beyond 461 stays hidden anyway)

def main():
    os.makedirs(RENDER_DIR, exist_ok=True)
    print("loading workbook (full mode, data_only)…", flush=True)
    wb = openpyxl.load_workbook(SRC, data_only=True)
    ws = wb["Boat Module"]
    # Drop the hidden Dropdowns sheet from the render copy.
    for name in list(wb.sheetnames):
        if name != "Boat Module":
            del wb[name]

    print("hiding rows/cols…", flush=True)
    keep_rows = {1, 2, 3, HIGHFIELD_DIVIDER_ROW, ROW}
    for r in range(1, ws.max_row + 1):
        if r not in keep_rows:
            ws.row_dimensions[r].hidden = True
    keep = set(KEEP_COLS)
    for c in range(1, max(LAST_COL, ws.max_column) + 1):
        letter = get_column_letter(c)
        if c not in keep:
            ws.column_dimensions[letter].hidden = True
        else:
            # readable widths for the exhibit
            ws.column_dimensions[letter].hidden = False
            ws.column_dimensions[letter].width = 28 if c in (3, 312, 313, 315, 390, 402, 403) else 14

    ws.print_area = f"A1:{get_column_letter(LAST_COL)}{ROW}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)

    print("saving render copy…", flush=True)
    wb.save(RENDER_XLSX)
    wb.close()

    print("converting to PDF…", flush=True)
    subprocess.run(
        ["soffice", "--headless", "-env:UserInstallation=file:///tmp/ultimate/lo-profile2",
         "--convert-to", "pdf", "--outdir", RENDER_DIR, RENDER_XLSX],
        check=True, timeout=550,
    )
    pdf = f"{RENDER_DIR}/mpf-quote.pdf"
    if not os.path.exists(pdf):
        sys.exit("PDF not produced")

    print("rasterising PNG…", flush=True)
    import pypdfium2 as pdfium
    doc = pdfium.PdfDocument(pdf)
    for i in range(min(len(doc), 2)):
        bmp = doc[i].render(scale=150 / 72)
        img = bmp.to_pil()
        suffix = "" if i == 0 else f"-p{i+1}"
        img.save(f"{RENDER_DIR}/mpf-quote{suffix}.png")
    doc.close()
    print("done:", os.listdir(RENDER_DIR))

if __name__ == "__main__":
    main()
