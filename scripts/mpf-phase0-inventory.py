#!/usr/bin/env python3
"""MPF Phase 0 — full inventory of NSM's Master Price File family.

Reads every workbook in tasks/mpf-source/ (read-only, cached values) and
produces a machine-readable inventory: every sheet, its dimensions, the
header row, sample values per column, defined names, and the external
cross-workbook link map (parsed straight from each xlsx's internal
xl/externalLinks/ XML — the "insane amount of links" made visible).

Outputs:
  tasks/mpf-audit/inventory.json   (machine evidence)
  tasks/mpf-audit/INVENTORY.md     (human summary)
Appends actions to tasks/mpf-audit/AUDIT_LOG.jsonl.
"""
import json, os, re, zipfile, datetime
from openpyxl import load_workbook

SRC = "tasks/mpf-source"
OUT = "tasks/mpf-audit"

def log(action, detail):
    with open(f"{OUT}/AUDIT_LOG.jsonl", "a") as f:
        f.write(json.dumps({"ts": datetime.datetime.now(datetime.timezone.utc).isoformat(),
                            "action": action, "detail": detail}) + "\n")

def external_links(path):
    """Cross-workbook links, straight from the xlsx internals."""
    targets = []
    try:
        with zipfile.ZipFile(path) as z:
            for n in z.namelist():
                if re.match(r"xl/externalLinks/_rels/externalLink\d+\.xml\.rels", n):
                    xml = z.read(n).decode("utf-8", "ignore")
                    targets += re.findall(r'Target="([^"]+)"', xml)
    except Exception as e:
        targets = [f"(unreadable: {e})"]
    return targets

inventory = {"generatedUtc": datetime.datetime.now(datetime.timezone.utc).isoformat(), "workbooks": []}

for fname in sorted(os.listdir(SRC)):
    path = os.path.join(SRC, fname)
    if fname.endswith(".zip"):
        continue
    entry = {"file": fname, "bytes": os.path.getsize(path)}
    if fname.lower().endswith(".csv"):
        with open(path, "r", errors="ignore") as f:
            lines = f.read().splitlines()
        entry.update({"type": "csv", "rows": len(lines),
                      "header": lines[0].split(",")[:40] if lines else []})
        inventory["workbooks"].append(entry)
        log("phase0.inventoried", {"file": fname, "type": "csv", "rows": len(lines)})
        print(f"  {fname}: csv, {len(lines)} rows")
        continue

    entry["type"] = "xlsx"
    entry["externalLinks"] = external_links(path)
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        entry["definedNames"] = sorted(wb.defined_names.keys())[:60] if hasattr(wb, "defined_names") else []
        sheets = []
        for ws in wb.worksheets:
            s = {"name": ws.title, "state": ws.sheet_state,
                 "maxRow": ws.max_row, "maxCol": ws.max_column}
            # header row + 2 sample rows (first 40 cols)
            rows = []
            for i, row in enumerate(ws.iter_rows(min_row=1, max_row=3, max_col=min(ws.max_column or 1, 40), values_only=True)):
                rows.append([str(v)[:48] if v is not None else None for v in row])
                if i >= 2:
                    break
            s["header"] = rows[0] if rows else []
            s["sample"] = rows[1:3]
            sheets.append(s)
        entry["sheets"] = sheets
        wb.close()
        log("phase0.inventoried", {"file": fname, "sheets": [(x['name'], x['maxRow'], x['maxCol']) for x in sheets],
                                   "externalLinks": entry["externalLinks"]})
        print(f"  {fname}: {len(sheets)} sheets, links={len(entry['externalLinks'])}")
    except Exception as e:
        entry["error"] = str(e)[:200]
        log("phase0.error", {"file": fname, "error": str(e)[:200]})
        print(f"  {fname}: ERROR {e}")
    inventory["workbooks"].append(entry)

json.dump(inventory, open(f"{OUT}/inventory.json", "w"), indent=1)

# human summary
with open(f"{OUT}/INVENTORY.md", "w") as f:
    f.write("# MPF Phase 0 — Inventory\n\n")
    f.write(f"Generated {inventory['generatedUtc']} from {len(inventory['workbooks'])} files in `tasks/mpf-source/` (source zip sha256 logged in AUDIT_LOG.jsonl).\n\n")
    f.write("| Workbook | Size | Sheets | Rows (max sheet) | External links |\n|---|---|---|---|---|\n")
    for w in inventory["workbooks"]:
        if w.get("type") == "csv":
            f.write(f"| {w['file']} | {w['bytes']:,} | (csv) | {w['rows']:,} | — |\n")
        else:
            sheets = w.get("sheets", [])
            maxrows = max((s["maxRow"] or 0) for s in sheets) if sheets else 0
            f.write(f"| {w['file']} | {w['bytes']:,} | {len(sheets)} | {maxrows:,} | {len(w.get('externalLinks', []))} |\n")
    f.write("\n## Sheets per workbook\n")
    for w in inventory["workbooks"]:
        if w.get("type") != "xlsx":
            continue
        f.write(f"\n### {w['file']}\n")
        for s in w.get("sheets", []):
            hdr = ", ".join([h for h in (s["header"] or []) if h][:12])
            f.write(f"- **{s['name']}** ({s['state']}) — {s['maxRow']:,} rows × {s['maxCol']} cols. Header: {hdr}\n")
        if w.get("externalLinks"):
            f.write(f"- External links: {'; '.join(w['externalLinks'][:10])}\n")
print("\nPhase 0 inventory written: tasks/mpf-audit/inventory.json + INVENTORY.md")
